import streamlit as st, pandas as pd, plotly.express as px

st.set_page_config(page_title="Dashboard TI", layout="wide")

# =========================================================================
# 🎬 TÍTULO ANIMADO FIXO NO TOPO DE TODAS AS ABAS
# =========================================================================
st.markdown(
    """
    <div style="background-color: #1F4E78; padding: 10px; border-radius: 5px; margin-bottom: 20px;">
        <marquee behavior="scroll" direction="left" scrollamount="7" style="color: #FFFFFF; font-size: 28px; font-weight: bold; font-family: Arial;">
            🚀 Villagemall Inventario
        </marquee>
    </div>
    """, 
    unsafe_allow_html=True
)

# =========================================================================
# 🌐 CONEXÃO COM O BANCO DE DADOS EM NUVEM (POSTGRESQL DO STREAMLIT)
# =========================================================================
conn = st.connection("postgresql", type="sql")

# Inicialização automática das tabelas na Nuvem
with conn.session as session:
    session.execute("CREATE TABLE IF NOT EXISTS servidores (nome TEXT, rack TEXT, ip TEXT, localidade TEXT);")
    session.execute("CREATE TABLE IF NOT EXISTS rede_ip (nome_computador TEXT, local TEXT, status TEXT, ip TEXT);")
    session.execute("CREATE TABLE IF NOT EXISTS pinpad (local TEXT, modelo TEXT, num_serie TEXT);")
    session.execute("CREATE TABLE IF NOT EXISTS ferramentas (nome TEXT, quantidade INTEGER, num_serie TEXT);")
    session.execute("CREATE TABLE IF NOT EXISTS bobinas (tamanho TEXT PRIMARY KEY, quantidade INTEGER DEFAULT 0);")
    session.execute("CREATE TABLE IF NOT EXISTS controle_ht (local TEXT PRIMARY KEY, quantidade INTEGER DEFAULT 0, status TEXT);")
    session.execute("CREATE TABLE IF NOT EXISTS chamados (id SERIAL PRIMARY KEY, num_chamado TEXT, empresa TEXT, status TEXT, descricao TEXT);")
    session.commit()

bobinas = ['58MMX300M ENTRADA', '76MMX300 PAGAMENTO', '79MM80 GUICHE', '57MMX08 PAPA FILA']
terminais = ["EL1", "EVIP R1", "EVIP R2", "ESS1 R1", "ESS1 R2", "ESS1 R3", "ESS2 R1", "ESS2 R2", "ESS2 R3",
             "PAGAMENTO PORTARIA C", "PAGAMENTO PORTARIA B", "VIP PANORAMICO", "SS1 AMERICAS", "SS1 PANORAMICO",
             "SS1 LAGOA1", "SS1 LAGOA2", "SS1 LAGOA3", "GUICHE SS1 LAGOA", "SS2 PANORAMICO", "TPC SS2 AMERICAS",
             "TPC SS2 PANORAMICO", "TPC SS2 LAGOA"]
hts = ["ACEESO SSS1 AMERICAS", "SVIP", "E3", "E4", "E5", "E6", "E8", "E9", "E10", "E1"]

# Popula tabelas base na nuvem se estiverem vazias
for b in bobinas:
    df_check = conn.query(f"SELECT 1 FROM bobinas WHERE tamanho = '{b}'", ttl=0)
    if df_check.empty:
        with conn.session as session:
            session.execute("INSERT INTO bobinas (tamanho, quantidade) VALUES (:t, 0);", {"t": b})
            session.commit()

for h in hts:
    df_check = conn.query(f"SELECT 1 FROM controle_ht WHERE local = '{h}'", ttl=0)
    if df_check.empty:
        with conn.session as session:
            session.execute("INSERT INTO controle_ht (local, quantidade, status) VALUES (:l, 0, 'Disponível');", {"l": h})
            session.commit()

menu = st.sidebar.radio("Menu:", ["📊 Dashboard", "🖥️ Servidores", "🌐 Rede/IP", "💳 Pinpad", "🔧 Ferramentas", "🧻 Bobinas", "📻 HT", "🎫 Chamados"])

# ----------------- 1. DASHBOARD -----------------
if menu == "📊 Dashboard":
    st.title("📊 Painel de Indicadores Gerais")
    
    total_servidores = len(conn.query("SELECT * FROM servidores", ttl=0))
    total_pinpads = len(conn.query("SELECT * FROM pinpad", ttl=0))
    
    ferr_df = conn.query("SELECT SUM(quantidade) as t FROM ferramentas", ttl=0)
    total_ferramentas = int(ferr_df['t'].fillna(0).iloc[0])
    
    bob_df = conn.query("SELECT SUM(quantidade) as t FROM bobinas", ttl=0)
    total_bobinas = int(bob_df['t'].fillna(0).iloc[0])
    
    ht_df = conn.query("SELECT SUM(quantidade) as t FROM controle_ht", ttl=0)
    total_hts = int(ht_df['t'].fillna(0).iloc[0])
    
    total_chamados = len(conn.query("SELECT * FROM chamados", ttl=0))

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Servidores", total_servidores)
    m2.metric("Pinpads", total_pinpads)
    m3.metric("Ferramentas (Qtd Total)", total_ferramentas)
    m4.metric("Chamados Registrados", total_chamados)
    st.write("---")

    def gerar_excel_customizado():
        import io
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import PatternFill, Font, Border, Side
        
        wb = Workbook()
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        thin_border = Border(left=Side(style='thin', color='BFBFBF'), right=Side(style='thin', color='BFBFBF'), top=Side(style='thin', color='BFBFBF'), bottom=Side(style='thin', color='BFBFBF'))
        
        # TAB 1: BOBINAS (Com Linhas de Grade e Gráfico)
        ws1 = wb.active; ws1.title = "Estoque de Bobinas"
        ws1.sheet_properties.tabColor = "1F4E78"; ws1.sheet_view.showGridLines = True
        df_b_excel = conn.query("SELECT tamanho, quantidade FROM bobinas", ttl=0)
        ws1.append(["Modelo de Bobina", "Quantidade"])
        for idx, row in df_b_excel.iterrows(): ws1.append([row['tamanho'], row['quantidade']])
        for row in ws1.iter_rows(min_row=1, max_row=len(df_b_excel)+1, min_col=1, max_col=2):
            for cell in row:
                cell.border = thin_border
                if cell.row == 1: cell.fill = header_fill; cell.font = header_font
        
        chart1 = BarChart(); chart1.type = "col"; chart1.title = "Gráfico de Barras: Estoque de Bobinas"
        chart1.add_data(Reference(ws1, min_col=2, min_row=1, max_row=len(df_b_excel)+1), titles_from_data=True)
        chart1.set_categories(Reference(ws1, min_col=1, min_row=2, max_row=len(df_b_excel)+1))
        ws1.add_chart(chart1, "D2")
        
        # RESTANTE DAS ORELHAS COLORIDAS
        abas = [
            ("Servidores", "2ECC71", "SELECT nome, rack, ip, localidade FROM servidores", ["Nome", "Rack", "IP", "Localidade"]),
            ("Pinpads", "9B59B6", "SELECT local, modelo, num_serie FROM pinpad", ["Local", "Modelo", "N° Série"]),
            ("Ferramentas", "F39C12", "SELECT nome, quantidade, num_serie FROM ferramentas", ["Nome", "Quantidade", "N° Série"]),
            ("Rádios HT", "34495E", "SELECT local, quantidade, status FROM controle_ht", ["Local", "Quantidade", "Status"]),
            ("Chamados", "16A085", "SELECT num_chamado, empresa, status, descricao FROM chamados", ["N° Chamado", "Empresa", "Status", "Descrição"])
        ]
        for title, color, query, headers in abas:
            ws = wb.create_sheet(title=title); ws.sheet_properties.tabColor = color; ws.sheet_view.showGridLines = True
            df = conn.query(query, ttl=0)
            ws.append(headers)
            for idx, row in df.iterrows(): ws.append(list(row))
            for row in ws.iter_rows(min_row=1, max_row=len(df)+1, min_col=1, max_col=len(headers)):
                for cell in row:
                    cell.border = thin_border
                    if cell.row == 1: cell.fill = header_fill; cell.font = header_font

        output = io.BytesIO(); wb.save(output); return output.getvalue()

    st.download_button(label="📥 Gerar Relatório Excel com Gráficos", data=gerar_excel_customizado(), file_name="Inventario_TI_Formatado.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="btn_excel_dashboard")
    st.write("---")

    col1, col2 = st.columns(2)
    with col1:
        st.subheader("📊 1. Gráfico de Barras: Detalhado de Bobinas")
        df_b = conn.query("SELECT * FROM bobinas", ttl=0)
        if df_b['quantidade'].sum() > 0:
            st.plotly_chart(px.bar(df_b, x='tamanho', y='quantidade', color='tamanho', labels={'quantidade': 'Qtd', 'tamanho': 'Modelo'}), use_container_width=True)
        else: st.info("Sem estoque de bobinas gravado.")
    with col2:
        st.subheader("🍕 2. Gráfico de Pizza: Distribuição Total do Inventário")
        total_redes = len(conn.query("SELECT * FROM rede_ip", ttl=0))
        dados_gerais = {
            "Item do Inventário": ["Servidores", "IPs de Rede", "Pinpads", "Ferramentas", "Bobinas em Estoque", "Rádios HT", "Chamados"],
            "Quantidade Gravada": [total_servidores, total_redes, total_pinpads, total_ferramentas, total_bobinas, total_hts, total_chamados]
        }
        df_inventario_total = pd.DataFrame(dados_gerais)
        if df_inventario_total["Quantidade Gravada"].sum() > 0:
            st.plotly_chart(px.pie(df_inventario_total, values="Quantidade Gravada", names="Item do Inventário", hole=0.15, color_discrete_sequence=px.colors.qualitative.Pastel), use_container_width=True)
        else: st.info("Nenhum dado gravado no sistema.")

# ----------------- 2. SERVIDORES -----------------
elif menu == "🖥️ Servidores":
    st.title("🖥️ Servidores")
    with st.form("f_serv", clear_on_submit=True):
        n, r, i, l = st.text_input("Nome"), st.text_input("Rack"), st.text_input("IP"), st.text_input("Localidade")
        if st.form_submit_button("Gravar") and n and i:
            with conn.session as session:
                session.execute("INSERT INTO servidores (nome, rack, ip, localidade) VALUES (:n, :r, :i, :l);", {"n":n, "r":r, "i":i, "l":l})
                session.commit()
            st.success("Salvo!"); st.rerun()
            
    df_serv = conn.query("SELECT * FROM servidores", ttl=0)
    st.dataframe(df_serv, use_container_width=True)
    if not df_serv.empty:
        st.write("---"); st.subheader("🗑️ Excluir Registro")
        serv_para_excluir = st.selectbox("Selecione o Servidor para remover:", df_serv['nome'].unique())
        if st.button("Excluir Servidor", type="primary"):
            with conn.session as session:
